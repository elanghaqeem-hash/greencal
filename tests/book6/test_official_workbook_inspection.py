import json
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Protection

from scripts.book6.inspect_official_workbook import inspect_workbook, validate_mapping


class OfficialWorkbookInspectionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / "ojk-uart-template.xlsx"
        w = Workbook()
        s = w.active
        s.title = "Identitas"
        s["A1"] = "SYNTHETIC-BANK-SECRET"
        s["B1"] = 10
        s["C3"] = None
        s["C3"].fill = PatternFill(patternType="solid", fgColor="0000FF")
        s["C3"].protection = Protection(locked=False)
        s["D3"] = "=SUM(B1:B1)"
        s.protection.sheet = True
        w.create_sheet("B. Emisi")
        w.save(self.path)
        self.source_hash = sha256(self.path.read_bytes()).hexdigest()

    def test_inspects_worksheet_without_modifying_source_or_exposing_values(self):
        report = inspect_workbook(self.path)
        self.assertEqual(report["template_sha256"], self.source_hash)
        self.assertEqual(report["sheet_order"], ["Identitas", "B. Emisi"])
        self.assertFalse(report["validated_official_template"])
        self.assertFalse(report["bank_production_approved"])
        self.assertFalse(report["contains_cell_values"])
        self.assertEqual(report["sheets"][0]["formula_count"], 1)
        self.assertEqual(report["sheets"][0]["unlocked_cell_addresses"], ["C3"])
        self.assertTrue(report["sheets"][0]["sheet_protection_enabled"])
        self.assertNotIn("SYNTHETIC-BANK-SECRET", json.dumps(report))
        self.assertNotIn("=SUM(", json.dumps(report))
        self.assertEqual(sha256(self.path.read_bytes()).hexdigest(), self.source_hash)

    def test_approved_mapping_rejects_locked_formula_wrong_hash_and_blue_unconfirmed(self):
        report = inspect_workbook(self.path)
        mapping = {"template_sha256": self.source_hash, "reviewed_by": "Bank Risk Reviewer",
                   "reviewed_at": "2026-10-31",
                   "entries": [{"key": "bank.code", "sheet": "Identitas", "cell": "C3",
                                "style_id": 1, "confirmed_blue_input": True}]}
        mapping["entries"][0]["style_id"] = report["sheets"][0]["style_profiles"][0]["style_id"]
        # Synthetic workbook may have multiple styles; read exact style by openpyxl.
        from openpyxl import load_workbook
        w = load_workbook(self.path)
        mapping["entries"][0]["style_id"] = w["Identitas"]["C3"].style_id
        w.close()
        self.assertTrue(validate_mapping(self.path, report, mapping)["mapping_structure_valid"])
        for key, value in [("template_sha256", "0" * 64), ("reviewed_by", "")]:
            invalid = deepcopy(mapping)
            invalid[key] = value
            with self.assertRaises(ValueError):
                validate_mapping(self.path, report, invalid)
        invalid = deepcopy(mapping)
        invalid["entries"][0]["confirmed_blue_input"] = False
        with self.assertRaisesRegex(ValueError, "blue-cell"):
            validate_mapping(self.path, report, invalid)
        invalid = deepcopy(mapping)
        invalid["entries"][0]["cell"] = "D3"
        invalid["entries"][0]["style_id"] = 0
        with self.assertRaisesRegex(ValueError, "formula or locked"):
            validate_mapping(self.path, report, invalid)

    def test_non_excel_or_encrypted_file_fails_closed_without_logging_secret(self):
        invalid = Path(self.tmp.name) / "bad.xlsx"
        invalid.write_bytes(b"not xlsx; not encrypted")
        with self.assertRaisesRegex(ValueError, "password"):
            inspect_workbook(invalid, password_env="BOOK6_SYNTHETIC_PASSWORD_UNSET")


if __name__ == "__main__":
    unittest.main()
