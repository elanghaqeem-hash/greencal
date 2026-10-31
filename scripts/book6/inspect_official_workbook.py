#!/usr/bin/env python3
"""Inspect an OJK Buku 6 Excel template without changing cells or disclosing values.

This tool is a controlled pre-release inspector, NOT a regulatory report writer.
It supports encrypted Office files via msoffcrypto-tool and in-file sheet
protection via openpyxl. Passwords are read exclusively from an environment
variable, never accepted on the command line or printed.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import re
import sys
from zipfile import is_zipfile

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string

MAX_CELL_ADDRESSES = 20000
MAX_WORKBOOK_BYTES = 20_000_000


def _color(c):
    if c is None:
        return None
    if c.type == "rgb":
        return {"type": "rgb", "value": str(c.rgb)}
    if c.type == "indexed":
        return {"type": "indexed", "value": str(c.indexed)}
    if c.type == "theme":
        return {"type": "theme", "value": str(c.theme), "tint": c.tint}
    return {"type": str(c.type), "value": None}


def _bytes(path, password_env):
    raw = Path(path).read_bytes()
    if not raw or len(raw) > MAX_WORKBOOK_BYTES:
        raise ValueError("Workbook is empty or exceeds 20 MB safety limit.")
    digest = sha256(raw).hexdigest()
    if is_zipfile(BytesIO(raw)):
        return raw, digest, False
    password = os.environ.get(password_env)
    if not password:
        raise ValueError("File may be Office-encrypted. Provide password via environment variable " + password_env + ".")
    try:
        import msoffcrypto
        office = msoffcrypto.OfficeFile(BytesIO(raw))
        office.load_key(password=password)
        target = BytesIO()
        office.decrypt(target)
        plain = target.getvalue()
    except Exception as exc:
        raise ValueError("Unable to decrypt Office workbook; check password and file integrity.") from exc
    if not is_zipfile(BytesIO(plain)):
        raise ValueError("Decrypted payload is not an XLSX/ZIP workbook.")
    return plain, digest, True


def inspect_workbook(path, password_env="BOOK6_XLSX_PASSWORD"):
    plain, source_hash, encrypted = _bytes(path, password_env)
    book = load_workbook(BytesIO(plain), read_only=False, data_only=False, keep_links=False)
    result = {
        "schema": "greencal-ojk-book6-template-inspection-v1",
        "template_sha256": source_hash,
        "encrypted_office_file": encrypted,
        "sheet_order": list(book.sheetnames),
        "worksheet_count": len(book.sheetnames),
        "workbook_structure_locked": bool(book.security.lockStructure) if book.security else False,
        "validated_official_template": False,
        "bank_production_approved": False,
        "contains_cell_values": False,
        "sheets": [],
    }
    for ws in book.worksheets:
        formula_fingerprint = sha256()
        formula_count = 0
        populated = 0
        unlocked = []
        styles = Counter()
        for cell in ws._cells.values():
            if cell.value is not None:
                populated += 1
            if cell.data_type == "f":
                formula_count += 1
                formula_fingerprint.update((cell.coordinate + ":" + str(cell.value) + "\n").encode())
            if not cell.protection.locked:
                unlocked.append(cell.coordinate)
            if cell.has_style:
                style_key = (
                    cell.style_id,
                    json.dumps(_color(cell.fill.fgColor), sort_keys=True),
                    json.dumps(_color(cell.font.color), sort_keys=True),
                    bool(cell.protection.locked),
                )
                styles[style_key] += 1
        if len(unlocked) > MAX_CELL_ADDRESSES:
            visible = unlocked[:MAX_CELL_ADDRESSES]
        else:
            visible = unlocked
        result["sheets"].append({
            "name": ws.title,
            "state": ws.sheet_state,
            "dimension": ws.calculate_dimension(),
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "sheet_protection_enabled": bool(ws.protection.sheet),
            "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
            "data_validation_count": len(ws.data_validations.dataValidation) if ws.data_validations else 0,
            "formula_count": formula_count,
            "formula_fingerprint_sha256": formula_fingerprint.hexdigest(),
            "populated_cell_count": populated,
            "unlocked_cell_count": len(unlocked),
            "unlocked_cell_addresses": visible,
            "unlocked_addresses_truncated": len(unlocked) > MAX_CELL_ADDRESSES,
            "style_profiles": [
                {"style_id": style_id, "fill_color": json.loads(fill),
                 "font_color": json.loads(font), "locked": locked, "count": count}
                for (style_id, fill, font, locked), count in sorted(styles.items())
            ],
        })
    book.close()
    return result


def validate_mapping(template_path, report, mapping, password_env="BOOK6_XLSX_PASSWORD"):
    """Review a separately approved address map; do not populate or save Excel."""
    if mapping.get("template_sha256") != report["template_sha256"]:
        raise ValueError("Approved map template hash does not match the exact XLSX file.")
    if not mapping.get("reviewed_by") or not mapping.get("reviewed_at"):
        raise ValueError("A named reviewer and review date are mandatory.")
    entries = mapping.get("entries")
    if not isinstance(entries, list) or not entries:
        raise ValueError("An explicit bank-reviewed cell map is mandatory.")
    plain, _, _ = _bytes(template_path, password_env)
    book = load_workbook(BytesIO(plain), data_only=False, read_only=False, keep_links=False)
    used_keys, used_cells = set(), set()
    for entry in entries:
        if not isinstance(entry, dict):
            raise ValueError("Invalid mapping entry.")
        key, sheet, addr = entry.get("key"), entry.get("sheet"), entry.get("cell")
        if not isinstance(key, str) or not re.fullmatch(r"[A-Za-z0-9_.-]{1,120}", key):
            raise ValueError("Invalid mapping key.")
        if sheet not in book.sheetnames or not isinstance(addr, str):
            raise ValueError("Mapped sheet/cell does not exist.")
        try:
            coordinate_from_string(addr)
        except (TypeError, ValueError) as exc:
            raise ValueError("Invalid cell address.") from exc
        ws = book[sheet]
        cell = ws[addr]
        if key in used_keys or (sheet, addr) in used_cells:
            raise ValueError("Duplicate key or target cell in mapping.")
        if cell.data_type == "f" or cell.protection.locked:
            raise ValueError("Mapping points to a formula or locked cell: " + sheet + "!" + addr)
        if entry.get("style_id") != cell.style_id:
            raise ValueError("Cell style differs from approved style: " + sheet + "!" + addr)
        if entry.get("confirmed_blue_input") is not True:
            raise ValueError("Each input cell requires explicit blue-cell visual confirmation.")
        if any(cell.coordinate in r for r in ws.merged_cells.ranges if cell.coordinate != r.start_cell.coordinate):
            raise ValueError("Cannot address non-anchor merged cell.")
        used_keys.add(key)
        used_cells.add((sheet, addr))
    book.close()
    return {"validated_entries": len(entries), "mapping_structure_valid": True,
            "official_regulatory_filing_validated": False, "bank_production_approved": False}


def main(argv=None):
    p = argparse.ArgumentParser(description="OJK Buku 6 encrypted/protected Excel read-only structure inspector")
    p.add_argument("--workbook", required=True, type=Path)
    p.add_argument("--output", required=True, type=Path)
    p.add_argument("--password-env", default="BOOK6_XLSX_PASSWORD")
    p.add_argument("--approved-map", type=Path, default=None)
    a = p.parse_args(argv)
    try:
        report = inspect_workbook(a.workbook, a.password_env)
        if a.approved_map:
            mapping = json.loads(a.approved_map.read_text(encoding="utf-8"))
            report["mapping_review"] = validate_mapping(a.workbook, report, mapping, a.password_env)
        a.output.parent.mkdir(parents=True, exist_ok=True)
        a.output.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        print("OJK Book 6 inspection complete; workbook unchanged; "
              + str(report["worksheet_count"]) + " sheets; mapped: "
              + str(report.get("mapping_review", {}).get("validated_entries", 0)) + ".")
    except Exception as exc:
        print("OJK Book 6 inspection failed: " + str(exc), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
